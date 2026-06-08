import tkinter as tk
from tkinter import ttk, filedialog, messagebox, simpledialog
import os
from docx import Document

from config import GRAMS_FILE_NAME, DEFAULTS_DIR

class NgramsEditorWindow:
    def __init__(self, parent, db_path, is_defaults=False):
        self.parent = parent
        self.db_path = db_path
        self.changes_made = False  # Флаг изменений
        self.is_defaults = is_defaults

        self.window = tk.Toplevel(parent)
        self.window.title("Список N-грамм")
        self.window.geometry("880x550")
        self.window.transient(parent)
        self.window.grab_set()

        text_frame = ttk.Frame(self.window)
        text_frame.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)

        self.listbox = tk.Listbox(text_frame)
        scrollbar_y = ttk.Scrollbar(text_frame, orient=tk.VERTICAL, command=self.listbox.yview)
        self.listbox.configure(yscrollcommand=scrollbar_y.set)
        self.listbox.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scrollbar_y.pack(side=tk.RIGHT, fill=tk.Y)

        self.grams_file_path = os.path.join(self.db_path, GRAMS_FILE_NAME)
        if os.path.exists(self.grams_file_path):
            with open(self.grams_file_path, 'r', encoding='utf-8') as f:
                grams = [line.strip() for line in f if line.strip()]
            for gram in sorted(grams):
                self.listbox.insert(tk.END, gram)

        self.original_grams = list(self.listbox.get(0, tk.END)) if self.listbox.size() > 0 else []
        self.ngrams_search_var = tk.StringVar()
        ttk.Label(self.window, text="Поиск N-грамм:").pack(anchor=tk.W, padx=10, pady=(5, 0))
        self.ngrams_search_entry = ttk.Entry(self.window, textvariable=self.ngrams_search_var)
        self.ngrams_search_entry.pack(fill=tk.X, padx=10, pady=(0, 5))
        self.ngrams_search_var.trace_add('write', lambda *args: self._search_ngrams())

        button_frame = ttk.Frame(self.window)
        button_frame.pack(fill=tk.X, padx=10, pady=5)

        self.add_btn = ttk.Button(button_frame, text="Добавить", command=self.add_gram)
        self.add_btn.pack(side=tk.LEFT, padx=5)

        ttk.Button(button_frame, text="Редактировать", command=self._edit_selected).pack(side=tk.LEFT, padx=5)

        self.del_btn = ttk.Button(button_frame, text="Удалить", command=self.delete_selected)
        self.del_btn.pack(side=tk.LEFT, padx=5)

        self.import_btn = ttk.Button(button_frame, text="Импортировать", command=self.import_from_file)
        self.import_btn.pack(side=tk.LEFT, padx=5)
        self.tooltip = None
        self.import_btn.bind("<Enter>",
                             lambda e: self.on_import_btn_hover("Каждая N-грамма должна быть с новой строки."))
        self.import_btn.bind("<Leave>", lambda e: self.hide_tooltip())

        if not self.is_defaults:
            ttk.Button(button_frame, text="Добавить значения по умолчанию",
                       command=self._merge_defaults).pack(side=tk.LEFT, padx=5)

        ttk.Button(button_frame, text="Удалить всё", command=self.delete_all).pack(side=tk.LEFT, padx=5)

        self.close_btn = ttk.Button(button_frame, text="Закрыть", command=self.on_close)
        self.close_btn.pack(side=tk.RIGHT, padx=5)

        self.window.bind('<Delete>',
                         lambda e: self.delete_selected() if self.listbox.focus_get() == self.listbox else None)

    def delete_all(self):
        if self.listbox.size() == 0: return
        dlg = tk.Toplevel(self.window)
        dlg.title("Удаление всех N-грамм")
        dlg.geometry("400x150")
        dlg.transient(self.window)
        dlg.grab_set()
        ttk.Label(dlg, text="Для удаления всех N-грамм введите \"УДАЛИТЬ\":").pack(pady=10)
        entry = ttk.Entry(dlg, width=30)
        entry.pack(pady=5)
        entry.focus()
        def execute():
            if entry.get().strip() == "УДАЛИТЬ":
                self.listbox.delete(0, tk.END)
                self.changes_made = True
                dlg.destroy()
                messagebox.showinfo("Успех", "Все N-граммы удалены.")
            else:
                messagebox.showwarning("Ошибка", "Вы не ввели \"УДАЛИТЬ\".")
        ttk.Button(dlg, text="Удалить", command=execute).pack(pady=10)
        ttk.Button(dlg, text="Отмена", command=dlg.destroy).pack(pady=2)

    def _search_ngrams(self):
        query = self.ngrams_search_var.get().lower()
        self.listbox.delete(0, tk.END)
        for gram in self.original_grams:
            if query in gram.lower():
                self.listbox.insert(tk.END, gram)

    def _edit_selected(self):
        sel = self.listbox.curselection()
        if not sel: return
        old = self.listbox.get(sel[0])
        new = simpledialog.askstring("Переименовать", "Новое название:", initialvalue=old, parent=self.window)
        if new and new.strip() and new != old:
            self.listbox.delete(sel[0]); self.listbox.insert(sel[0], new.strip())
            self.listbox.selection_set(sel[0]); self.changes_made = True

    def _merge_defaults(self):
        if not messagebox.askyesno("Подтверждение", "Вы уверены, что хотите добавить значения по умолчанию?"): return

        def_path = os.path.join(DEFAULTS_DIR, "grams.txt")  # для фракций: factions.txt и т.д.
        if not os.path.exists(def_path): return

        current = set(self.listbox.get(0, tk.END))
        added = 0
        with open(def_path, 'r', encoding='utf-8') as f:
            for line in f:
                item = line.strip()
                if item and item not in current:
                    self.listbox.insert(tk.END, item)
                    self.changes_made = True
                    added += 1
        messagebox.showinfo("Готово", f"Добавлено {added} новых значений из параметров по умолчанию.")

    def on_import_btn_hover(self, text):
        self.hide_tooltip()
        x, y, _, _ = self.import_btn.bbox("insert")
        x += self.import_btn.winfo_rootx() + 25
        y += self.import_btn.winfo_rooty() + 25
        self.tooltip = tk.Toplevel(self.window)
        self.tooltip.wm_overrideredirect(True)
        self.tooltip.wm_geometry(f"+{x}+{y}")
        label = ttk.Label(self.tooltip, text=text, background="#ffffe0", relief="solid", borderwidth=1,
                          font=("tahoma", "8", "normal"))
        label.pack(ipadx=1, ipady=1)

    def hide_tooltip(self):
        if self.tooltip:
            self.tooltip.destroy()
            self.tooltip = None

    def add_gram(self):
        new_gram = simpledialog.askstring("Добавить N-грамму", "Введите n-грамму:", parent=self.window)
        if new_gram:
            new_gram = new_gram.strip()
            if new_gram:
                existing = self.listbox.get(0, tk.END)
                if new_gram in existing:
                    messagebox.showwarning("Предупреждение", f"N-грамма '{new_gram}' уже существует.")
                    return
                self.listbox.insert(tk.END, new_gram)
                self.changes_made = True  # Устанавливаем флаг при добавлении
                self.original_grams = [g for g in self.listbox.get(0, tk.END)]

    def delete_selected(self):
        selection = self.listbox.curselection()
        idx = selection[0]
        if selection:
            self.listbox.delete(selection[0])
            self.changes_made = True  # Устанавливаем флаг при удалении

            if self.listbox.size() > 0:
                new_idx = idx if idx < self.listbox.size() else idx - 1
                self.listbox.select_set(new_idx)
                self.listbox.activate(new_idx)

            self.original_grams = [g for g in self.listbox.get(0, tk.END)]

    def import_from_file(self):
        file_path = filedialog.askopenfilename(
            title="Выберите файл с N-граммами",
            filetypes=[("Text files", "*.txt"), ("Word documents", "*.docx"), ("All files", "*.*")]
        )
        if file_path:
            new_grams = set()
            try:
                if file_path.lower().endswith('.docx'):
                    doc = Document(file_path)
                    full_text = '\n'.join([para.text for para in doc.paragraphs])
                else:
                    with open(file_path, 'r', encoding='utf-8') as f:
                        full_text = f.read()

                imported = [line.strip() for line in full_text.split('\n') if line.strip()]

                current_grams = set(self.listbox.get(0, tk.END))
                for gram in imported:
                    if gram and gram not in current_grams:
                        new_grams.add(gram)

                for gram in sorted(new_grams):
                    self.listbox.insert(tk.END, gram)

                if new_grams:
                    self.changes_made = True  # <-- ИСПРАВЛЕНИЕ: Устанавливаем флаг при импорте
                    messagebox.showinfo("Импорт", f"Добавлено {len(new_grams)} новых N-грамм.")
                    self.original_grams = [g for g in self.listbox.get(0, tk.END)]
                else:
                    messagebox.showinfo("Импорт", "Новых N-грамм не найдено или все уже существуют.")
            except Exception as e:
                messagebox.showerror("Ошибка", f"Не удалось импортировать файл: {str(e)}")

    def on_close(self):
        if self.changes_made:
            if messagebox.askyesno("Сохранить изменения?", "Вы хотите сохранить изменения в списке N-грамм?"):
                self.save_ngrams()
        self.window.destroy()

    def save_ngrams(self):
        try:
            with open(self.grams_file_path, 'w', encoding='utf-8') as f:
                grams = self.listbox.get(0, tk.END)
                for gram in grams:
                    f.write(gram + '\n')
            os.utime(self.db_path, None)
            messagebox.showinfo("Сохранение", "Список N-грамм сохранён.")
            self.changes_made = False  # Сбрасываем флаг после сохранения
        except Exception as e:
            messagebox.showerror("Ошибка", f"Не удалось сохранить файл: {str(e)}")